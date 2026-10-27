import io
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document as DocxWriter
from langchain_core.language_models.fake_chat_models import FakeListChatModel
from rest_framework.test import APIClient

from . import rag_service
from .models import Document, QAHistory

MEDIA = tempfile.mkdtemp()


def make_docx(*paragraphs, table=None) -> bytes:
    doc = DocxWriter()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(client, name, data):
    return client.post(
        "/api/documents/",
        {"file": SimpleUploadedFile(name, data, content_type="application/octet-stream")},
        format="multipart",
    )


@override_settings(MEDIA_ROOT=MEDIA, RAG_DEMO_MODE=True)
class RagApiTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.client = APIClient()
        rag_service.reset_index_cache()

    # ── documents ───────────────────────────────────────────────────────
    def test_upload_extracts_paragraphs_and_tables(self):
        data = make_docx("قرارداد پشتیبانی نرم‌افزار", table=[["مبلغ", "۱۰۰ میلیون"]])
        res = upload(self.client, "contract.docx", data)
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.data["title"], "contract")  # defaults to file name
        self.assertEqual(res.data["status"], "ready")
        doc = Document.objects.get(pk=res.data["id"])
        self.assertIn("قرارداد پشتیبانی", doc.full_text)
        self.assertIn("مبلغ | ۱۰۰ میلیون", doc.full_text)

    def test_rejects_non_docx_and_fake_docx(self):
        self.assertEqual(upload(self.client, "a.txt", b"hello").status_code, 400)
        self.assertEqual(upload(self.client, "fake.docx", b"not a zip").status_code, 400)
        self.assertEqual(Document.objects.count(), 0)

    def test_empty_docx_is_marked_failed_without_polluting_text(self):
        res = upload(self.client, "empty.docx", make_docx())
        self.assertEqual(res.status_code, 201)
        self.assertEqual(res.data["status"], "failed")
        self.assertTrue(res.data["error_message"])
        self.assertEqual(Document.objects.get(pk=res.data["id"]).full_text, "")

    @override_settings(MAX_UPLOAD_MB=0)
    def test_size_limit(self):
        self.assertEqual(upload(self.client, "big.docx", make_docx("x")).status_code, 400)

    def test_delete_removes_file_from_disk(self):
        res = upload(self.client, "gone.docx", make_docx("متن"))
        doc = Document.objects.get(pk=res.data["id"])
        path = doc.file.path
        self.assertEqual(self.client.delete(f"/api/documents/{doc.id}/").status_code, 204)
        import os
        self.assertFalse(os.path.exists(path))

    def test_list_is_lightweight_detail_has_full_text(self):
        res = upload(self.client, "d.docx", make_docx("سلام دنیا"))
        listing = self.client.get("/api/documents/").data
        self.assertNotIn("full_text", listing[0])
        detail = self.client.get(f"/api/documents/{res.data['id']}/").data
        self.assertIn("سلام دنیا", detail["full_text"])

    # ── ask ─────────────────────────────────────────────────────────────
    def test_ask_without_documents(self):
        res = self.client.post("/api/ask/", {"question": "چیزی هست؟"}, format="json")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["sources"], [])
        self.assertIsNone(res.data["history_id"])

    def test_ask_validation(self):
        self.assertEqual(self.client.post("/api/ask/", {"question": "؟"}, format="json").status_code, 400)

    def test_ask_returns_sources_and_saves_history(self):
        upload(self.client, "sla.docx", make_docx(
            "زمان پاسخ‌گویی به درخواست‌های بحرانی حداکثر چهار ساعت کاری است.",
            "هزینه پشتیبانی سالانه معادل ده درصد مبلغ قرارداد خواهد بود.",
        ))
        res = self.client.post("/api/ask/", {"question": "زمان پاسخ‌گویی به درخواست بحرانی چقدر است؟"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.data["sources"][0]["title"], "sla")
        self.assertIn("چهار ساعت", res.data["answer"])
        self.assertEqual(QAHistory.objects.count(), 1)
        self.assertEqual(QAHistory.objects.get().sources[0]["doc_id"], res.data["sources"][0]["doc_id"])

    def test_ask_can_be_scoped_to_one_document(self):
        a = upload(self.client, "a.docx", make_docx("گربه‌ها حیوانات خانگی محبوبی هستند و شیر می‌نوشند."))
        b = upload(self.client, "b.docx", make_docx("پایتون یک زبان برنامه‌نویسی همه‌منظوره است."))
        res = self.client.post(
            "/api/ask/", {"question": "گربه چه می‌نوشد؟", "document_ids": [b.data["id"]]}, format="json"
        )
        self.assertEqual({s["doc_id"] for s in res.data["sources"]}, {b.data["id"]})
        self.assertNotEqual(a.data["id"], b.data["id"])

    def test_index_is_cached_until_documents_change(self):
        upload(self.client, "a.docx", make_docx("متن اول برای آزمایش کش نمایه."))
        calls = []
        original = rag_service.FAISS.from_texts
        rag_service.FAISS.from_texts = classmethod(lambda cls, *a, **k: calls.append(1) or original(*a, **k))
        try:
            for _ in range(3):
                self.client.post("/api/ask/", {"question": "کش نمایه چیست؟"}, format="json")
            self.assertEqual(len(calls), 1)
            upload(self.client, "b.docx", make_docx("سند دوم هم اضافه شد."))
            self.client.post("/api/ask/", {"question": "سند دوم چیست؟"}, format="json")
            self.assertEqual(len(calls), 2)
        finally:
            rag_service.FAISS.from_texts = original

    @override_settings(RAG_DEMO_MODE=False, OPENROUTER_API_KEY="")
    def test_missing_api_key_is_503_not_500(self):
        Document.objects.create(title="t", file=SimpleUploadedFile("t.docx", make_docx("متن آزمایشی")))
        res = self.client.post("/api/ask/", {"question": "سوال آزمایشی"}, format="json")
        self.assertEqual(res.status_code, 503)
        self.assertEqual(res.data["code"], "not_configured")

    @override_settings(RAG_DEMO_MODE=False, OPENROUTER_API_KEY="test-key")
    def test_llm_path_uses_retrieved_context(self):
        from unittest import mock

        upload_res = None
        with mock.patch.object(rag_service, "get_embeddings", rag_service.HashingEmbeddings):
            upload(self.client, "x.docx", make_docx("مرخصی سالانه کارکنان بیست و شش روز است."))
            with mock.patch.object(rag_service, "get_llm", lambda: FakeListChatModel(responses=["بیست و شش روز."])):
                res = self.client.post("/api/ask/", {"question": "مرخصی سالانه چند روز است؟"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.data["answer"], "بیست و شش روز.")
        self.assertIn("مرخصی سالانه", res.data["context_retrieved"])

    # ── history & status ────────────────────────────────────────────────
    def test_history_list_delete_clear_and_status(self):
        for i in range(2):
            QAHistory.objects.create(question=f"q{i}", answer="a", sources=[])
        page = self.client.get("/api/history/").data
        self.assertEqual(page["count"], 2)
        self.assertEqual(self.client.delete(f"/api/history/{page['results'][0]['id']}/").status_code, 204)
        self.assertEqual(self.client.delete("/api/history/clear/").data["deleted"], 1)
        st = self.client.get("/api/status/").data
        self.assertTrue(st["configured"] and st["demo_mode"])
        self.assertEqual(st["questions"], 0)

    def test_admin_is_gone(self):
        self.assertEqual(self.client.get("/admin/").status_code, 404)

    def test_seed_demo_command(self):
        from django.core.management import call_command

        call_command("seed_demo", stdout=io.StringIO())
        self.assertEqual(Document.objects.filter(status="ready").count(), 3)
        self.assertEqual(QAHistory.objects.count(), 5)
        self.assertTrue(all(h.sources for h in QAHistory.objects.all()))
        call_command("seed_demo", stdout=io.StringIO())  # idempotent
        self.assertEqual(Document.objects.count(), 3)
